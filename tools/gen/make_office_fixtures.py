# -*- coding: utf-8 -*-
"""make_office_fixtures.py — Office-файлы без Word: tests/fixtures/office/{таблица.xlsx, слайды.pptx, письмо.odt}.

Что (все с кириллицей и картинкой из tests/fixtures/images):
  таблица.xlsx — лист «Города»: шапка, 5 строк, формула СУММ, формат чисел, картинка chart.png;
  слайды.pptx  — 3 слайда: заголовок, список, слайд с картинкой photo.jpg;
  письмо.odt   — заголовок, абзацы, картинка logo.png (ODF 1.2, odfpy).
  docx/rtf делает tools/gen/make_office_fixtures.ps1 (Word COM).
Откуда: openpyxl, python-pptx, odfpy, Pillow из venv C:\\Tools\\venv-test.
Детерминизм: свойства документа (автор, даты) — константы; ZIP пересобирается с датой 1980-01-01
и тем же порядком записей (normalize_word.rezip); в ODT «mimetype» первым и без сжатия.
Запуск: "C:\\Tools\\venv-test\\Scripts\\python.exe" -X utf8 -I tools/gen/make_office_fixtures.py [--out <корень>]
"""
import re
import sys
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True                     # без __pycache__ в tools/gen
sys.path.insert(0, str(Path(__file__).resolve().parent))

import _genlib as g  # noqa: E402
import _texts as tx  # noqa: E402
import normalize_word as nw  # noqa: E402

FIXED = g.FIXED_DT.replace(tzinfo=None)


def images_dir(root: Path) -> Path:
    d = root / "tests" / "fixtures" / "images"
    return d if (d / "photo.jpg").exists() else g.REPO / "tests" / "fixtures" / "images"


def make_xlsx(dst: Path, img: Path, tmp: Path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Города"
    ws.append(["Город", "Население", "Регион", "Доля, %"])
    rows = [("Москва", 13100000, "Москва"), ("Казань", 1300000, "Татарстан"), ("Сочи", 470000, "Краснодарский край"),
            ("Тверь", 410000, "Тверская обл."), ("Ёлки", 1200, "Посёлок")]
    for i, (c, n, r) in enumerate(rows, start=2):
        ws.append([c, n, r, None])
        ws.cell(row=i, column=4, value="=B%d/B$8*100" % i).number_format = "0.00"
        ws.cell(row=i, column=2).number_format = "# ##0"
    ws["A8"] = "Итого"
    ws["B8"] = "=SUM(B2:B6)"
    ws["B8"].number_format = "# ##0"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7F5")
        c.alignment = Alignment(horizontal="center")
    ws["A8"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    pic = XLImage(str(img))
    pic.width, pic.height = 400, 250
    ws.add_image(pic, "F2")
    ws2 = wb.create_sheet("Заметки")
    for i, s in enumerate(tx.RU[:5], start=1):
        ws2.cell(row=i, column=1, value=s)
    p = wb.properties
    p.creator = nw.AUTHOR_ASCII
    p.lastModifiedBy = nw.AUTHOR_ASCII
    p.title = "Таблица городов"
    p.created = FIXED
    p.modified = FIXED
    raw = tmp / "table.raw.xlsx"
    wb.save(raw)
    nw.norm_ooxml(raw, dst)


def make_pptx(dst: Path, img: Path, tmp: Path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Оттиск: пробные слайды"
    s1.placeholders[1].text = "Кириллица, латиница и картинка"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Что умеет программа"
    body = s2.placeholders[1].text_frame
    body.text = tx.RU[0]
    for t in tx.RU[1:4]:
        para = body.add_paragraph()
        para.text = t
        para.font.size = Pt(18)
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Картинка"
    s3.shapes.add_picture(str(img), Inches(1.5), Inches(1.8), width=Inches(7))
    cp = prs.core_properties
    cp.author = nw.AUTHOR_ASCII
    cp.last_modified_by = nw.AUTHOR_ASCII
    cp.title = "Пробные слайды"
    cp.created = FIXED
    cp.modified = FIXED
    cp.revision = 1
    raw = tmp / "slides.raw.pptx"
    prs.save(raw)
    nw.norm_ooxml(raw, dst)


def make_odt(dst: Path, img: Path, tmp: Path):
    from odf import draw, text
    from odf.opendocument import OpenDocumentText
    from odf.style import GraphicProperties, ParagraphProperties, Style, TextProperties
    doc = OpenDocumentText()
    h = Style(name="Heading", family="paragraph")
    h.addElement(TextProperties(fontsize="18pt", fontweight="bold", fontfamily="Arial"))
    doc.styles.addElement(h)
    body = Style(name="Body", family="paragraph")
    body.addElement(TextProperties(fontsize="12pt", fontfamily="Times New Roman"))
    body.addElement(ParagraphProperties(margintop="0.1cm", marginbottom="0.2cm"))
    doc.styles.addElement(body)
    fr = Style(name="fr1", family="graphic")
    fr.addElement(GraphicProperties(anchortype="paragraph"))
    doc.automaticstyles.addElement(fr)
    doc.text.addElement(text.H(outlinelevel=1, stylename=h, text="Письмо"))
    doc.text.addElement(text.P(stylename=body, text="Здравствуйте!"))
    for par in tx.RU[:3]:
        doc.text.addElement(text.P(stylename=body, text=par))
    doc.text.addElement(text.P(stylename=body, text="Mit freundlichen Grüßen — с уважением, Оттиск."))
    href = doc.addPicture("Pictures/logo.png", "image/png", img.read_bytes())   # без uuid в имени
    p = text.P(stylename=body)
    frame = draw.Frame(stylename=fr, width="3cm", height="3cm", anchortype="paragraph", name="logo")
    frame.addElement(draw.Image(href=href))
    p.addElement(frame)
    doc.text.addElement(p)
    raw = tmp / "letter.raw.odt"
    doc.save(str(raw))
    # нормализация: meta.xml (даты, генератор с версией Python), ZIP без дат
    entries = []
    with zipfile.ZipFile(raw) as z:
        names = z.namelist()
        order = ["mimetype"] + [n for n in names if n != "mimetype"]
        for n in order:
            data = z.read(n)
            if n == "meta.xml":
                t = data.decode("utf-8")
                t = re.sub(r"<meta:creation-date>[^<]*</meta:creation-date>",
                           "<meta:creation-date>2026-01-01T12:00:00</meta:creation-date>", t)
                t = re.sub(r"<dc:date>[^<]*</dc:date>", "<dc:date>2026-01-01T12:00:00</dc:date>", t)
                t = re.sub(r"<meta:generator>[^<]*</meta:generator>",
                           "<meta:generator>odfpy (Ottisk tools/gen)</meta:generator>", t)
                t = re.sub(r"<dc:creator>[^<]*</dc:creator>", "<dc:creator>%s</dc:creator>" % nw.AUTHOR_ASCII, t)
                data = t.encode("utf-8")
            entries.append((n, data, zipfile.ZIP_STORED if n == "mimetype" else zipfile.ZIP_DEFLATED))
    nw.rezip(entries, dst)


def main():
    args = g.common_args(__doc__)
    root = Path(args.out)
    dst = root / "tests" / "fixtures" / "office"
    dst.mkdir(parents=True, exist_ok=True)
    tmp = g.tmp_dir(args, "office-py")
    imgs = images_dir(root)
    make_xlsx(dst / "таблица.xlsx", imgs / "chart.png", tmp)
    make_pptx(dst / "слайды.pptx", imgs / "photo.jpg", tmp)
    make_odt(dst / "письмо.odt", imgs / "logo.png", tmp)
    for n in ("таблица.xlsx", "слайды.pptx", "письмо.odt"):
        with zipfile.ZipFile(dst / n) as z:
            bad = z.testzip()
            if bad:
                raise RuntimeError("%s: битая запись %s" % (n, bad))
        g.say("office: %-14s %7d байт" % (n, (dst / n).stat().st_size))


if __name__ == "__main__":
    main()
